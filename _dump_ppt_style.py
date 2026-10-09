"""只读：抽取参考 deck 的版式基因（画布/主题色/字体/每页填充与字号）。"""
import collections
import sys
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = sys.argv[1] if len(sys.argv) > 1 else (
    r"C:\Users\Ma Xuancheng\Documents\xwechat_files\wxid_so6xr9pwdl522_f026\msg\file"
    r"\2026-09\雷达前沿大会2026-多模态电磁大模型及其应用研究-史治国-浙江大学.pptx")

z = zipfile.ZipFile(SRC)
print("=== zip 条目（theme/layout 数）")
names = z.namelist()
print("slides:", sum(1 for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")))
print("layouts:", sum(1 for n in names if n.startswith("ppt/slideLayouts/layout") and n.endswith(".xml")))
print("masters:", sum(1 for n in names if n.startswith("ppt/slideMasters/") and n.endswith(".xml")))
print("media:", sum(1 for n in names if n.startswith("ppt/media/")))

import re  # noqa: E402
t = z.read("ppt/theme/theme1.xml").decode("utf-8", "replace")
clr = re.search(r"<a:clrScheme.*?</a:clrScheme>", t, re.S)
print("\n=== clrScheme")
for m in re.finditer(r'<a:(\w+)>\s*<a:(?:srgbClr val="([0-9A-Fa-f]{6})"|sysClr[^>]*lastClr="([0-9A-Fa-f]{6})")',
                     clr.group(0) if clr else ""):
    print(f"  {m.group(1):10s} #{m.group(2) or m.group(3)}")
fs = re.search(r"<a:fontScheme.*?</a:fontScheme>", t, re.S)
if fs:
    for kind in ("majorFont", "minorFont"):
        m = re.search(rf"<a:{kind}>(.*?)</a:{kind}>", fs.group(0), re.S)
        if m:
            lat = re.search(r'<a:latin typeface="([^"]*)"', m.group(1))
            ea = re.search(r'<a:ea typeface="([^"]*)"', m.group(1))
            print(f"  {kind}: latin={lat.group(1) if lat else '-'} ea={ea.group(1) if ea else '-'}")

prs = Presentation(SRC)
print(f"\n=== 画布 {prs.slide_width} x {prs.slide_height} EMU "
      f"= {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")

fills = collections.Counter()
sizes = collections.Counter()
fonts = collections.Counter()
for i, s in enumerate(prs.slides, 1):
    page_fill, page_txt = [], []
    for sh in s.shapes:
        try:
            f = sh.fill
            if f.type is not None and f.type == 1:
                page_fill.append(str(f.fore_color.rgb))
        except Exception:
            pass
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    sz = r.font.size.pt if r.font.size else None
                    col = None
                    try:
                        col = str(r.font.color.rgb) if r.font.color and r.font.color.type == 1 else None
                    except Exception:
                        pass
                    page_txt.append((sz, col, bool(r.font.bold), r.font.name, r.text.strip()[:24]))
                    if sz:
                        sizes[round(sz)] += 1
                    if r.font.name:
                        fonts[r.font.name] += 1
                    if col:
                        fills["txt:" + col] += 1
    for c in page_fill:
        fills["fill:" + c] += 1
    print(f"\n--- P{i:02d} 形状 {len(s.shapes)}  填充 {sorted(set(page_fill))}")
    for sz, col, b, fn, txt in page_txt[:6]:
        print(f"    {sz if sz else '-':>5} {col or '-':>8} {'B' if b else ' '} {fn or '-':<16} {txt}")

print("\n=== 全局 填充/文字色 top")
for k, v in fills.most_common(18):
    print(f"  {k:16s} {v}")
print("=== 字号 top")
for k, v in sorted(sizes.items()):
    print(f"  {k:>4} pt : {v}")
print("=== 字体 top")
for k, v in fonts.most_common(10):
    print(f"  {k:24s} {v}")

Path("_ppt_ref_style.txt").write_text("", encoding="utf-8")
